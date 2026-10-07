# -*- coding: utf-8 -*-
"""SCM 지정출고 주간작업 — 순수 로직 (streamlit 비의존, 테스트 가능).

요약서 파싱 / 재고 파싱 / ② 재고 LOT 조인 / 우선순위 / 엑셀 작업지시서.
② 재고 매칭 규칙: '요약서 제조일자와 동일한 LOT'의 출고가능 재고만 집계.
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent.parent))
from core.shelf_life import parse_ymd, _today_kst  # noqa: E402


# ---------------- 파싱 유틸 ----------------
def _norm(s) -> str:
    return re.sub(r"\s+", "", str(s)).replace("★", "")


def _ymd8(value) -> str:
    d = parse_ymd(value)
    return d.strftime("%Y%m%d") if d else ""


def _md(value) -> str:
    d = parse_ymd(value)
    return d.strftime("%m/%d") if d else ""


def _clean_code(v) -> str:
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


_SUMMARY_MAP = [
    ("gubun", ["구분"]),
    ("code", ["품목코드"]),
    ("name", ["품목명"]),
    ("mfg", ["제조일자"]),
    ("exp", ["소비기한"]),
    ("track", ["기준트랙", "트랙"]),
    ("reach", ["도달일자"]),
    ("end", ["종료일자"]),
    ("partners", ["지정/해제거래처", "거래처"]),
    ("warehouse", ["보관창고"]),
    ("reason", ["처리사유"]),
    ("avail_box", ["가용재고(Box)", "가용재고"]),
    ("move_box", ["입고요청(Box)", "입고요청"]),
]


def parse_summary(file_bytes: bytes) -> pd.DataFrame:
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    hdr_idx = None
    for i, r in enumerate(rows):
        norms = [_norm(c) for c in r if c is not None]
        if "구분" in norms and "품목코드" in norms:
            hdr_idx = i
            break
    if hdr_idx is None:
        raise ValueError("요약서에서 헤더(구분/품목코드) 행을 찾지 못했습니다.")

    header = rows[hdr_idx]
    col = {}
    for ci, cell in enumerate(header):
        if cell is None:
            continue
        n = _norm(cell)
        for key, pats in _SUMMARY_MAP:
            if key in col:
                continue
            if any(_norm(p) in n for p in pats):
                col[key] = ci
                break
    if "code" not in col or "gubun" not in col:
        raise ValueError("요약서 헤더에서 '구분' 또는 '품목코드' 열을 찾지 못했습니다.")

    def get(r, key):
        ci = col.get(key)
        if ci is None or ci >= len(r):
            return None
        return r[ci]

    recs = []
    last_gubun = None
    for r in rows[hdr_idx + 1:]:
        code = get(r, "code")
        if code is None or _norm(code) == "품목코드":
            continue
        code_s = _clean_code(code)
        if not re.fullmatch(r"\d{3,}", code_s):
            continue
        gubun = get(r, "gubun")
        if gubun is not None and str(gubun).strip():
            last_gubun = str(gubun).strip()
        recs.append({
            "구분": last_gubun or "",
            "품목코드": code_s,
            "품목명": get(r, "name") or "",
            "제조일자": _ymd8(get(r, "mfg")),
            "소비기한": _ymd8(get(r, "exp")),
            "기준트랙": get(r, "track") or "",
            "도달일자": parse_ymd(get(r, "reach")),
            "종료일자": parse_ymd(get(r, "end")),
            "거래처": get(r, "partners") or "",
            "보관창고": (str(get(r, "warehouse")).strip() if get(r, "warehouse") is not None else ""),
            "처리사유": get(r, "reason") or "",
            "가용Box": get(r, "avail_box"),
            "요청Box": get(r, "move_box"),
        })
    return pd.DataFrame(recs)


def _pick_col(df: pd.DataFrame, *patterns):
    norms = {c: _norm(c) for c in df.columns}
    for pat in patterns:  # 완전일치 우선
        p = _norm(pat)
        for c, n in norms.items():
            if n == p:
                return c
    for pat in patterns:  # 포함검색
        p = _norm(pat)
        for c, n in norms.items():
            if p in n:
                return c
    return None


def parse_inventory(file_bytes: bytes) -> pd.DataFrame:
    df = pd.read_excel(io.BytesIO(file_bytes))
    c_code = _pick_col(df, "제품코드", "품목코드")
    c_mfg = _pick_col(df, "제조일")
    c_exp = _pick_col(df, "소비기한")
    c_loc = _pick_col(df, "Location")
    c_wh = _pick_col(df, "Inventory")
    c_cur = _pick_col(df, "현재고(Box)")
    c_avail = _pick_col(df, "출고가능(Box)")
    if not c_code or not c_avail:
        raise ValueError("재고 파일에서 '제품코드' 또는 '출고가능(Box)' 열을 찾지 못했습니다.")

    out = pd.DataFrame({
        "code": df[c_code].apply(_clean_code),
        "mfg": df[c_mfg].apply(_ymd8) if c_mfg else "",
        "창고": df[c_wh].astype(str).str.strip() if c_wh else "",
        "Location": df[c_loc].astype(str).str.strip() if c_loc else "",
        "소비기한": df[c_exp].apply(_ymd8) if c_exp else "",
        "현재고Box": pd.to_numeric(df[c_cur], errors="coerce").fillna(0).astype(int) if c_cur else 0,
        "출고가능Box": pd.to_numeric(df[c_avail], errors="coerce").fillna(0).astype(int),
    })
    out = out[out["code"].str.fullmatch(r"\d{3,}")]
    return out.reset_index(drop=True)


# ---------------- 우선순위 ----------------
def priority_label(reach_date) -> str:
    if reach_date is None:
        return ""
    d = (reach_date - _today_kst()).days
    if d < 0:
        return f"🔴 지남 ({-d}일)"
    if d <= 7:
        return f"🟠 이번주 (D-{d})"
    if d <= 14:
        return f"🟡 D-{d}"
    return f"D-{d}"


# ---------------- ② 재고 조인 ----------------
def join_move_stock(move_df: pd.DataFrame, inv) -> pd.DataFrame:
    rows = []
    for _, it in move_df.iterrows():
        code, mfg = it["품목코드"], it["제조일자"]
        req = it["요청Box"]
        req = int(req) if pd.notna(req) and str(req) != "" else None
        lot_avail = None
        lot_locs = ""
        total_avail = None
        note = ""
        if inv is not None:
            sub = inv[inv["code"] == code]
            total_avail = int(sub["출고가능Box"].sum()) if not sub.empty else 0
            lot = sub[(sub["mfg"] == mfg) & (sub["출고가능Box"] > 0)]
            lot_avail = int(lot["출고가능Box"].sum()) if not lot.empty else 0
            lot_locs = ", ".join(f"{r.Location}({r.출고가능Box})" for r in lot.itertuples())
            if lot_avail == 0:
                note = "⚠️ 해당 LOT 출고가능 재고 없음"
                if total_avail and total_avail > 0:
                    note += f" (다른 LOT 총 {total_avail}Box)"
            elif req is not None and lot_avail < req:
                note = f"⚠️ 부족 (요청 {req} > LOT가능 {lot_avail})"
            else:
                note = "✅ 충분"
        rows.append({
            "우선순위": priority_label(it["도달일자"]),
            "품목코드": code,
            "품목명": it["품목명"],
            "보관창고(→이동처)": it["보관창고"],
            "요청Box": req,
            "LOT(제조일)": mfg,
            "LOT출고가능Box": lot_avail,
            "집을 로케이션": lot_locs,
            "전체LOT출고가능": total_avail,
            "상태": note,
            "도달일자": _md(it["도달일자"]),
            "종료일자": _md(it["종료일자"]),
            "처리사유": it["처리사유"],
            "완료": False,
        })
    return pd.DataFrame(rows)


def simple_view(df: pd.DataFrame) -> pd.DataFrame:
    base = []
    for _, it in df.iterrows():
        base.append({
            "우선순위": priority_label(it["도달일자"]),
            "품목코드": it["품목코드"],
            "품목명": it["품목명"],
            "보관창고": it["보관창고"],
            "거래처": it["거래처"],
            "가용Box": it["가용Box"],
            "도달일자": _md(it["도달일자"]),
            "종료일자": _md(it["종료일자"]),
            "처리사유": it["처리사유"],
            "완료": False,
        })
    return pd.DataFrame(base)


# ---------------- 엑셀 작업지시서 ----------------
_FILL_HDR = PatternFill("solid", fgColor="305496")
_FILL_WARN = PatternFill("solid", fgColor="FFCDD2")
_FILL_OK = PatternFill("solid", fgColor="E8F5E9")


def build_xlsx(groups: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, df in groups.items():
        ws = wb.create_sheet(title=sheet_name[:31])
        if df is None or df.empty:
            ws.append(["(해당 없음)"])
            continue
        cols = list(df.columns)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _FILL_HDR
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for _, r in df.iterrows():
            ws.append([("" if pd.isna(v) else v) for v in r.tolist()])
            status = str(r.get("상태", ""))
            if status.startswith("⚠️"):
                for c in ws[ws.max_row]:
                    c.fill = _FILL_WARN
            elif status.startswith("✅"):
                for c in ws[ws.max_row]:
                    c.fill = _FILL_OK
        for ci, cn in enumerate(cols, start=1):
            w = 14
            if cn in ("품목명", "집을 로케이션", "처리사유", "거래처"):
                w = 34
            elif cn in ("상태", "우선순위", "보관창고(→이동처)"):
                w = 18
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A2"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
